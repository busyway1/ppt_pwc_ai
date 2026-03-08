#!/usr/bin/env python3
"""
Classify a shape group segment into a component type.

Uses spatial heuristics (zone, shape count, size patterns) and text content
analysis against component-types.json keywords.

Module API:
    classify_component(segment, component_types) -> str
"""

import argparse
import json
import re
import sys
from pathlib import Path

# Sequential text patterns (numbers, steps)
SEQUENTIAL_PATTERNS = [
    re.compile(r"^\d+[\.\)]?\s"),  # "1. " or "1) "
    re.compile(r"^step\s*\d", re.IGNORECASE),
    re.compile(r"^\d+단계"),
    re.compile(r"^phase\s*\d", re.IGNORECASE),
]


def _get_all_text(shapes):
    """Collect all text from a list of shapes."""
    texts = []
    for shape in shapes:
        paragraphs = shape.paragraphs if hasattr(shape, "paragraphs") else []
        for para in paragraphs:
            if para.text.strip():
                texts.append(para.text.strip())
    return texts


def _get_max_font_size(shapes):
    """Get the maximum font size across all shapes."""
    max_size = 0
    for shape in shapes:
        paragraphs = shape.paragraphs if hasattr(shape, "paragraphs") else []
        for para in paragraphs:
            if para.font_size and para.font_size > max_size:
                max_size = para.font_size
    return max_size


def _get_min_font_size(shapes):
    """Get the minimum font size across all shapes (ignoring None/0)."""
    min_size = 999
    for shape in shapes:
        paragraphs = shape.paragraphs if hasattr(shape, "paragraphs") else []
        for para in paragraphs:
            if para.font_size and 0 < para.font_size < min_size:
                min_size = para.font_size
    return min_size if min_size < 999 else 0


def _has_sequential_text(texts):
    """Check if texts contain sequential numbering patterns."""
    matches = 0
    for text in texts:
        for pat in SEQUENTIAL_PATTERNS:
            if pat.search(text):
                matches += 1
                break
    return matches >= 2


def _has_connectors_or_arrows(shapes):
    """Check if shapes include connectors or arrow-like shapes."""
    for shape in shapes:
        shape_obj = shape.shape if hasattr(shape, "shape") else None
        if shape_obj is None:
            continue
        # Check shape type name for connector/arrow indicators
        if hasattr(shape_obj, "shape_type"):
            st = str(shape_obj.shape_type)
            if "CONNECTOR" in st.upper() or "ARROW" in st.upper():
                return True
        # Check auto shape type name in XML
        if hasattr(shape_obj, "element"):
            elem_xml = (
                shape_obj.element.xml if hasattr(shape_obj.element, "xml") else ""
            )
            if "cxnSp" in elem_xml or "arrow" in elem_xml.lower():
                return True
    return False


def _has_chart(shapes):
    """Check if any shape is a chart."""
    for shape in shapes:
        shape_obj = shape.shape if hasattr(shape, "shape") else None
        if shape_obj is None:
            continue
        if hasattr(shape_obj, "has_chart") and shape_obj.has_chart:
            return True
    return False


def _has_table(shapes):
    """Check if any shape is a table."""
    for shape in shapes:
        shape_obj = shape.shape if hasattr(shape, "shape") else None
        if shape_obj is None:
            continue
        if hasattr(shape_obj, "has_table") and shape_obj.has_table:
            return True
    return False


def _has_images_with_text(shapes):
    """Check if shapes include images paired with text."""
    has_image = False
    has_text = False
    for shape in shapes:
        shape_obj = shape.shape if hasattr(shape, "shape") else None
        if shape_obj and hasattr(shape_obj, "image"):
            try:
                _ = shape_obj.image
                has_image = True
            except Exception:
                pass
        paragraphs = shape.paragraphs if hasattr(shape, "paragraphs") else []
        if any(p.text.strip() for p in paragraphs):
            has_text = True
    return has_image and has_text


def _shapes_similar_size(shapes, tolerance=0.15):
    """Check if multiple shapes are similarly sized."""
    if len(shapes) < 2:
        return False
    ref = shapes[0]
    count = 1
    for s in shapes[1:]:
        if ref.width == 0 or ref.height == 0:
            continue
        w_ratio = abs(s.width - ref.width) / max(s.width, ref.width)
        h_ratio = abs(s.height - ref.height) / max(s.height, ref.height)
        if w_ratio <= tolerance and h_ratio <= tolerance:
            count += 1
    return count >= 3


def _is_two_column(shapes, slide_width):
    """Check if shapes form a two-column layout."""
    if len(shapes) < 2:
        return False

    # Find two large shapes that each take ~half the width
    large_shapes = [s for s in shapes if s.width > slide_width * 0.3]
    if len(large_shapes) < 2:
        return False

    # Sort by left position
    large_shapes.sort(key=lambda s: s.left)
    left_shape = large_shapes[0]
    right_shape = large_shapes[-1]

    # Check roughly equal widths
    if left_shape.width == 0:
        return False
    ratio = abs(left_shape.width - right_shape.width) / max(
        left_shape.width, right_shape.width
    )
    if ratio > 0.3:
        return False

    # Check they don't overlap horizontally too much
    overlap = min(
        left_shape.left + left_shape.width, right_shape.left + right_shape.width
    ) - max(left_shape.left, right_shape.left)
    if overlap > left_shape.width * 0.3:
        return False

    return True


def _covers_most_area(shapes, slide_width, slide_height):
    """Check if shapes collectively cover >80% of slide area."""
    slide_area = slide_width * slide_height
    if slide_area == 0:
        return False

    total_area = sum(s.width * s.height for s in shapes)
    return (total_area / slide_area) > 0.8


def _keyword_score(texts, component_types):
    """Score each component type based on keyword matches in text.

    Returns dict of type_id -> match_count.
    """
    scores = {}
    combined = " ".join(texts).lower()
    for ct in component_types:
        count = 0
        for kw in ct.get("keywords", []):
            if kw.lower() in combined:
                count += 1
        if count > 0:
            scores[ct["id"]] = count
    return scores


def classify_component(segment, component_types, slide_width=None, slide_height=None):
    """Classify a segment into a component type.

    Args:
        segment: dict with zone, shapes, bounding_box (from segment_slide)
        component_types: list of type dicts from component-types.json
        slide_width: slide width in inches (optional, for two-column/full-page detection)
        slide_height: slide height in inches (optional)

    Returns:
        str: component type id (e.g. "card-grid", "header")
    """
    # LLM이 분류한 타입이 있으면 바로 반환 (18타입 지원)
    llm_type = segment.get("llm_type")
    if llm_type:
        return llm_type

    zone = segment.get("zone", "body")
    shapes = segment.get("shapes", [])
    is_grid = segment.get("is_grid", False)

    if not shapes:
        return "card-grid"

    texts = _get_all_text(shapes)
    max_font = _get_max_font_size(shapes)
    min_font = _get_min_font_size(shapes)

    # Decoration zone
    if zone == "decoration":
        paragraphs = shapes[0].paragraphs if hasattr(shapes[0], "paragraphs") else []
        total_text = "".join(p.text for p in paragraphs).strip()
        if len(total_text) <= 2:
            return "accent-bar"

    # Header zone: zone 자체가 header이면 높은 확률로 header
    # 폰트 크기가 명시적일 때는 추가 검증, 없으면 zone 기반 판단
    if zone == "header" and len(shapes) <= 5:
        if max_font == 0 or max_font >= 18:
            return "header"

    # Footer zone: zone 기반 판단 (폰트 크기 없어도 footer로 분류)
    if zone == "footer":
        if min_font == 0 or min_font <= 12:
            return "footer"

    # Chart detection
    if _has_chart(shapes):
        return "chart-area"

    # Table detection
    if _has_table(shapes):
        return "table-layout"

    # Full-page (covers >80% of slide)
    if (
        slide_width
        and slide_height
        and _covers_most_area(shapes, slide_width, slide_height)
    ):
        if len(shapes) <= 5:
            return "full-page"

    # Two-column detection
    if slide_width and _is_two_column(shapes, slide_width):
        return "two-column"

    # Process-flow: connectors/arrows between shapes
    if _has_connectors_or_arrows(shapes):
        return "process-flow"

    # Grid-flagged from segmentation: card-grid or timeline
    if is_grid or _shapes_similar_size(shapes):
        if _has_sequential_text(texts):
            return "timeline"
        return "card-grid"

    # Sequential text with linear arrangement
    if _has_sequential_text(texts) and len(shapes) >= 3:
        return "timeline"

    # Images paired with text
    if _has_images_with_text(shapes):
        return "icon-grid"

    # Keyword-based secondary signal
    kw_scores = _keyword_score(texts, component_types)
    if kw_scores:
        best_type = max(kw_scores, key=kw_scores.get)
        return best_type

    # Default fallback
    return "card-grid"


def main():
    """CLI demo: classify all segments in a PPTX file."""
    parser = argparse.ArgumentParser(
        description="Classify slide segments into component types."
    )
    parser.add_argument("input", help="Input PowerPoint file (.pptx)")
    args = parser.parse_args()

    PROJECT_ROOT = Path(__file__).resolve().parent.parent
    sys.path.insert(0, str(PROJECT_ROOT / "tools"))
    sys.path.insert(0, str(PROJECT_ROOT / "scripts"))

    from inventory import extract_text_inventory, ShapeData
    from segment_slide import segment_slide
    from pptx import Presentation

    pptx_path = Path(args.input)
    if not pptx_path.exists():
        print(f"Error: File not found: {pptx_path}")
        sys.exit(1)

    # Load component types
    config_path = PROJECT_ROOT / "config" / "component-types.json"
    with open(config_path, "r", encoding="utf-8") as f:
        component_types = json.load(f)["types"]

    prs = Presentation(str(pptx_path))
    slide_width = round(ShapeData.emu_to_inches(prs.slide_width), 2)
    slide_height = round(ShapeData.emu_to_inches(prs.slide_height), 2)

    inventory = extract_text_inventory(pptx_path, prs)

    for slide_id, shapes_dict in inventory.items():
        shapes_list = list(shapes_dict.values())
        segments = segment_slide(shapes_list, slide_width, slide_height)

        print(f"\n=== {slide_id} ===")
        for i, seg in enumerate(segments):
            comp_type = classify_component(
                seg, component_types, slide_width, slide_height
            )
            n = len(seg["shapes"])
            print(f"  Segment {i}: {comp_type} (zone={seg['zone']}, shapes={n})")


if __name__ == "__main__":
    main()
